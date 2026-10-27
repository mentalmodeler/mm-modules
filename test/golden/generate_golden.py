"""Writes test/fixtures/pyfcm-golden.json from PyFCM's scenario math.

Usage: python3 test/golden/generate_golden.py <path to PyFCM's Adjacency_Matrix_Example.xlsx>
Needs numpy and openpyxl. The JS tests only read the JSON; they never run Python.
"""
import datetime
import itertools
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from pyfcm_scenario import TooSlow, run_scenario  # noqa: E402

RULES = {"k": "kosko", "mk": "modified kosko", "r": "rescaled kosko"}
SQUASH = {"sig": "sigmoid", "tanh": "hyperbolic tangent", "biv": "bivalent", "triv": "trivalent"}


def read_matrix(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    n = ws.max_row - 1
    cell = lambda r, c: 0 if ws.cell(r + 1, c + 1).value is None else ws.cell(r + 1, c + 1).value
    names = [cell(0, j) for j in range(1, n + 1)]
    matrix = [[float(cell(i, j)) for j in range(1, n + 1)] for i in range(1, n + 1)]
    return names, matrix


# mm-modules stops after this many steps (src/inference.js), PyFCM never does.
MAX_ITERATIONS = 10000


def make_case(names, matrix, rule, squash, lam, noise, activate):
    """Returns None for a case PyFCM does not settle within MAX_ITERATIONS: PyFCM would carry on
    (or never finish), mm-modules gives up at the cap, so the two are meant to differ there."""
    params = {"noise_threshold": noise, "infer_rule": rule, "function_type": squash,
              "lambda": lam, "activate": activate, "max_iterations": MAX_ITERATIONS}
    try:
        changes = run_scenario(matrix, names, params)["changes"]
    except TooSlow:
        return None
    return {
        "settings": {"rule": RULES[rule], "squash": SQUASH[squash], "lambda": lam, "noiseThreshold": noise},
        "clamps": [{"name": k, "value": v} for k, v in activate.items()],
        "changes": changes,
    }


# Small maps with weights sitting exactly on the noise thresholds (0.25, 0.5) and feedback loops
# of both signs, so the bivalent/trivalent squash functions and the threshold's inclusive boundary
# (abs(w) <= threshold is dropped) are checked against PyFCM, not just the large example map.
SMALL_MAPS = [
    {"name": "branching with a negative loop", "names": ["A", "B", "C", "D", "E"], "edges": [
        ("A", "B", 0.6), ("A", "C", -0.4), ("B", "D", 0.8), ("C", "D", 0.5), ("D", "E", -0.7), ("E", "C", 0.3)]},
    {"name": "positive loop with an inhibitor", "names": ["A", "B", "C", "D"], "edges": [
        ("A", "B", 0.5), ("B", "C", 0.5), ("C", "A", 0.5), ("D", "C", -0.25)]},
    {"name": "two-concept negative loop on the thresholds", "names": ["A", "B", "C"], "edges": [
        ("A", "B", 1.0), ("B", "A", -0.5), ("C", "A", 0.25)]},
]


def small_map(spec):
    names = spec["names"]
    matrix = [[0.0] * len(names) for _ in names]
    for source, target, weight in spec["edges"]:
        matrix[names.index(source)][names.index(target)] = weight
    return names, matrix


def main(path):
    names, matrix = read_matrix(path)
    cases = []
    activate = {names[1]: 0.5, names[4]: 1.0}
    for rule, squash, lam, noise in itertools.product(RULES, SQUASH, [1, 2.5], [0, 0.15, 0.25, 0.5]):
        cases.append(make_case(names, matrix, rule, squash, lam, noise, dict(activate)))
    # bivalent/trivalent only move when a clamp pushes a concept to 0 or below, so clamp to those
    for rule, squash, noise in itertools.product(RULES, ["biv", "triv"], [0, 0.25, 0.5]):
        for clamp in ({names[1]: 0.0}, {names[4]: -1.0, names[8]: 1.0}):
            cases.append(make_case(names, matrix, rule, squash, 1, noise, dict(clamp)))
    # an explicit clamp to exactly 0 (PyFCM allows it; the old runScenario cannot express it)
    for rule in ("k", "mk"):
        cases.append(make_case(names, matrix, rule, "sig", 1, 0, {names[1]: 0.5, names[2]: 0.0}))
    cases = [c for c in cases if c is not None]

    maps = []
    for spec in SMALL_MAPS:
        small_names, small_matrix = small_map(spec)
        small_cases = []
        for rule, squash, lam, noise in itertools.product(RULES, SQUASH, [1, 2], [0, 0.25, 0.5]):
            for clamp in ({"A": 1.0}, {"B": -0.5, "C": 0.0}):
                small_cases.append(make_case(small_names, small_matrix, rule, squash, lam, noise, dict(clamp)))
        small_cases = [c for c in small_cases if c is not None]
        maps.append({"name": spec["name"], "names": small_names, "matrix": small_matrix, "cases": small_cases})

    out = os.path.join(HERE, "..", "fixtures", "pyfcm-golden.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w") as f:
        json.dump({
            "source": "PyFCM %s %s" % (datetime.date.today().isoformat(), os.path.basename(path)),
            "names": names, "matrix": matrix, "cases": cases, "maps": maps,
        }, f)
    print("wrote", os.path.normpath(out), len(cases), "cases;",
          ", ".join("%s: %d" % (m["name"], len(m["cases"])) for m in maps))


if __name__ == "__main__":
    main(sys.argv[1])
